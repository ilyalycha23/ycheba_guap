# -*- coding: utf-8 -*-
"""Оформление отчётов ГУАП (титульный лист, стили, рисунки, таблицы)."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor, Emu

FONT = "Times New Roman"


def set_run_font(run, size=14, bold=False, italic=False, all_caps=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    if all_caps:
        run.font.all_caps = True


def set_paragraph_format(
    p,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line=1.25,
    space_before=0,
    space_after=0,
    line=1.5,
    left=0,
    right=0,
):
    pf = p.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Cm(first_line) if first_line else Cm(0)
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.left_indent = Cm(left)
    pf.right_indent = Cm(right)


def add_text(doc, text, *, size=14, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=1.25, space_before=0, space_after=0, line=1.5):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=align, first_line=first_line, space_before=space_before, space_after=space_after, line=line)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_gost(doc, text):
    return add_text(
        doc,
        text,
        size=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=0,
        space_before=12,
        space_after=6,
    )


def add_caption(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, space_before=6, space_after=12, line=1.0)
    run = p.add_run(text)
    set_run_font(run, size=12)
    return p


def add_picture(doc, path, width_cm=16.0):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, space_before=6, space_after=0, line=1.0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    return p


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), kwargs.get("val", "single"))
        element.set(qn("w:sz"), kwargs.get("sz", "4"))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), kwargs.get("color", "000000"))
        tcBorders.append(element)
    tcPr.append(tcBorders)


def shade_cell(cell, color="D9E2F3"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, *, bold=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    pf.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, h, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(cell)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if c == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(table.rows[r].cells[c], str(val), size=11, align=align)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_empty_lines(doc, n):
    for _ in range(n):
        p = doc.add_paragraph()
        set_paragraph_format(p, first_line=0, line=1.0)
        run = p.add_run("")
        set_run_font(run, size=14)


def make_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return doc


def add_title_page(doc, lab_no, lab_title, course="Корпоративные информационные системы"):
    """Титульный лист по образцу отчётов ГУАП."""
    add_text(doc, "ГУАП", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, bold=True)
    add_text(doc, "КАФЕДРА № 42", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, bold=True, space_after=18)
    add_empty_lines(doc, 2)
    add_text(doc, "ОТЧЕТ", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, bold=True)
    add_text(doc, "ЗАЩИЩЕН С ОЦЕНКОЙ", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0)
    add_empty_lines(doc, 1)
    add_text(doc, "ПРЕПОДАВАТЕЛЬ", align=WD_ALIGN_PARAGRAPH.LEFT, first_line=0, bold=True)
    add_text(
        doc,
        "должность, уч. степень, звание          подпись, дата          инициалы, фамилия",
        size=12,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line=0,
        italic=True,
        line=1.0,
    )
    add_empty_lines(doc, 3)
    add_text(
        doc,
        f"ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ №{lab_no}",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=0,
        bold=True,
    )
    add_text(doc, lab_title.upper(), align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, bold=True)
    add_text(doc, f"по курсу: {course}", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, italic=True)
    add_empty_lines(doc, 4)
    add_text(doc, "РАБОТУ ВЫПОЛНИЛ", align=WD_ALIGN_PARAGRAPH.LEFT, first_line=0, bold=True)
    add_text(doc, "СТУДЕНТ гр. № 4321                              И.А. Лаврешин", align=WD_ALIGN_PARAGRAPH.LEFT, first_line=0)
    add_text(
        doc,
        "подпись, дата                                          инициалы, фамилия",
        size=12,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line=0,
        italic=True,
        line=1.0,
    )
    add_empty_lines(doc, 4)
    add_text(doc, "Санкт-Петербург 2026", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0)
    doc.add_page_break()
