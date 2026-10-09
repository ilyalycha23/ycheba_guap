"""Generate Lab 1 Word report for linear regression."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

LAB1 = Path(__file__).resolve().parent
FIG = LAB1 / "figures"
RESULTS = json.loads((LAB1 / "results.json").read_text(encoding="utf-8"))


def set_run_font(run, name="Times New Roman", size=14, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_para(doc, text, *, size=14, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        pf.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_center(doc, text, *, size=14, bold=False, space_after=0):
    return add_para(
        doc, text, size=size, bold=bold, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=space_after
    )


def add_heading_num(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Cm(0)
    set_run_font(p.add_run(text), size=14, bold=True)


def add_picture_caption(doc, image_path: Path, caption: str, width_cm: float = 15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(8)
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_after = Pt(10)
    set_run_font(cap.add_run(caption), size=12)


def fmt(x: float, d: int = 3) -> str:
    if abs(x) >= 1000:
        return f"{x:,.{d}f}".replace(",", " ")
    return f"{x:.{d}f}"


def comma(x: float, d: int = 3) -> str:
    return fmt(x, d).replace(".", ",")


def build_title_page(doc: Document) -> None:
    items = [
        ("ГУАП", True, 14, 0),
        ("КАФЕДРА № 42", True, 14, 18),
        ("ОТЧЕТ ЗАЩИЩЕН С ОЦЕНКОЙ", True, 14, 6),
        ("ПРЕПОДАВАТЕЛЬ", False, 14, 6),
        ("д-р техн. наук, профессор", False, 14, 18),
        ("Фомин В.В.", True, 14, 6),
        ("должность, уч. степень, звание", False, 12, 12),
        ("подпись, дата", False, 12, 6),
        ("инициалы, фамилия", False, 12, 24),
        ("ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 1", True, 14, 6),
        ("Метод линейной регрессии", True, 14, 6),
        ("по курсу: Методы искусственного интеллекта", False, 14, 30),
        ("РАБОТУ ВЫПОЛНИЛ", False, 14, 6),
        ("СТУДЕНТ гр. №", False, 14, 6),
        ("4321", True, 14, 18),
        ("И.А. Лаврешин", True, 14, 12),
        ("подпись, дата", False, 12, 6),
        ("инициалы, фамилия", False, 12, 36),
        ("Санкт-Петербург 2026", False, 14, 0),
    ]
    for text, bold, size, after in items:
        add_center(doc, text, bold=bold, size=size, space_after=after)
    doc.add_page_break()


def build_body(doc: Document) -> None:
    m = RESULTS["metrics"]

    add_heading_num(doc, "1 Цель работы")
    add_para(
        doc,
        "Целью работы является освоение технологии построения модели простой линейной регрессии "
        "в среде KNIME Analytics Platform: загрузка табличных данных, разделение выборки на обучающую "
        "и тестовую части, обучение модели, получение прогнозов и оценка качества по метрикам Numeric Scorer, "
        "а также визуализация зависимости на диаграмме рассеяния.",
    )

    add_heading_num(doc, "2 Краткие теоретические сведения")
    add_para(
        doc,
        "Линейная регрессия — метод восстановления зависимости числового отклика Y от признака (признаков) X. "
        "В простейшем случае используется модель одной независимой переменной:",
    )
    add_center(doc, "Salarŷ = a + b · YearsExperience,", space_after=8)
    add_para(
        doc,
        f"где a — свободный член (Intercept), b — коэффициент при стаже работы YearsExperience, "
        f"Salarŷ — предсказанное значение заработной платы. Параметры a и b оцениваются по обучающей выборке "
        f"методом наименьших квадратов. В эксперименте получены оценки a ≈ {comma(RESULTS['intercept'], 2)}, "
        f"b ≈ {comma(RESULTS['coef_YearsExperience'], 2)}. Качество на тестовой выборке оценивается метриками "
        f"R², MAE, MSE, RMSE, MAPE и др.",
    )
    add_para(
        doc,
        "R² (коэффициент детерминации) показывает долю дисперсии целевой переменной, объяснённую моделью. "
        "MAE и RMSE характеризуют среднюю абсолютную и среднеквадратичную ошибку в единицах зарплаты; "
        "MAPE — среднюю относительную ошибку.",
    )

    add_heading_num(doc, "3 Описание исходных данных")
    add_para(
        doc,
        f"В работе использован учебный набор данных о зависимости заработной платы сотрудников от стажа работы "
        f"(файл Salary_Data.csv). Набор содержит {RESULTS['n_observations']} наблюдений и атрибуты: идентификатор "
        f"сотрудника (Employee), стаж работы в годах (YearsExperience) и заработную плату (Salary). "
        f"Целевая переменная — Salary; независимая переменная — YearsExperience.",
    )
    add_para(
        doc,
        "Файл размещён в каталоге lab1/data/Salary_Data.csv. Workflow собран по методическим указаниям "
        "(работа № 2 практикума KNIME); расчёты и рисунки воспроизведены с параметрами разбиения "
        "20 % Linear, как в методичке.",
    )

    add_heading_num(doc, "4 Ход выполнения работы")
    add_para(
        doc,
        "В соответствии с методическими указаниями собран workflow линейной регрессии. Состав узлов:",
    )
    for b in [
        "CSV Reader — загрузка файла Salary_Data.csv;",
        "Color Manager — цветовая маркировка строк по значению Salary (для визуализации);",
        "Partitioning — разделение данных: 20 % в обучающую выборку (порт 1), 80 % в тестовую (порт 2), способ выборки Linear;",
        "Linear Regression Learner — обучение модели, целевой атрибут Salary, независимый — YearsExperience;",
        "Regression Predictor — применение модели к тестовой выборке, формирование колонки Prediction (Salary);",
        "Numeric Scorer — расчёт метрик качества по парам «факт Salary — прогноз Prediction (Salary)»;",
        "Scatter Plot — диаграмма рассеяния YearsExperience–Salary.",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        set_run_font(p.add_run("– " + b), size=14)

    add_para(
        doc,
        f"Связи узлов: обучающая выборка ({RESULTS['train_size']} наблюдений, индексы {RESULTS['train_indices']}) "
        f"подаётся на Linear Regression Learner; тестовая выборка ({RESULTS['test_size']} наблюдений) и обученная "
        f"модель — на Regression Predictor; выход Predictor подключён к Numeric Scorer и Scatter Plot.",
    )
    add_picture_caption(
        doc,
        FIG / "fig1_workflow.png",
        "Рисунок 1 – Workflow линейной регрессии в KNIME Analytics Platform",
        16,
    )
    add_para(
        doc,
        "В настройках Numeric Scorer в качестве опорного (reference) столбца указан Salary, в качестве "
        "прогнозируемого (predicted) — Prediction (Salary).",
    )

    add_heading_num(doc, "5 Результаты эксперимента")
    add_para(
        doc,
        "После выполнения эксперимента получены прогнозы заработной платы на тестовой выборке и сводные метрики "
        "качества модели. Значения метрик приведены на рисунке 2 и в таблице 1.",
    )
    add_picture_caption(
        doc,
        FIG / "fig2_metrics.png",
        "Рисунок 2 – Метрики качества модели (узел Numeric Scorer, Lab1_Linear_Regression)",
        15,
    )
    add_center(doc, "Таблица 1 – Метрики качества линейной регрессии на тестовой выборке", size=12, space_after=6)

    table = doc.add_table(rows=8, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Метрика", "Значение"),
        ("R²", comma(m["R2"])),
        ("Mean Absolute Error (MAE)", comma(m["MAE"])),
        ("Mean Squared Error (MSE)", comma(m["MSE"])),
        ("Root Mean Squared Error (RMSE)", comma(m["RMSE"])),
        ("Mean Signed Difference", comma(m["MSD"])),
        ("Mean Absolute Percentage Error (MAPE)", f"{comma(m['MAPE'])} ({m['MAPE']*100:.1f} %)".replace(".", ",")),
        ("Adjusted R²", comma(m["AdjR2"])),
    ]
    for i, (a, b) in enumerate(rows):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_font(run, size=12, bold=(i == 0))

    doc.add_paragraph()
    add_para(
        doc,
        f"Полученное значение R² ≈ {comma(m['R2'])} означает, что модель объясняет около {m['R2']*100:.1f} % "
        f"дисперсии заработной платы на тестовой выборке за счёт линейной зависимости от стажа. Средняя абсолютная "
        f"ошибка составляет около {fmt(m['MAE']/1000, 1)} тыс. единиц зарплаты, RMSE ≈ {fmt(m['RMSE']/1000, 1)} тыс.; "
        f"относительная ошибка MAPE ≈ {m['MAPE']*100:.1f} %. Средняя знаковая ошибка равна {comma(m['MSD'])}. "
        f"Близость R² и adjusted R² объясняется использованием одного предиктора.",
    )
    add_picture_caption(doc, FIG / "fig4_predictions.png", "Рисунок 3 – Факт и прогноз Salary на тестовой выборке", 15.5)
    add_para(
        doc,
        "Диаграмма рассеяния (рисунок 4) подтверждает наличие выраженной положительной линейной связи "
        "между YearsExperience и Salary: с ростом стажа заработная плата увеличивается.",
    )
    add_picture_caption(
        doc,
        FIG / "fig3_scatter.png",
        "Рисунок 4 – Диаграмма рассеяния Salary от YearsExperience (результат workflow KNIME)",
        15,
    )
    add_para(
        doc,
        "Таким образом, экспериментально подтверждена применимость простой линейной регрессии для "
        "прогнозирования заработной платы по стажу на рассматриваемом наборе данных.",
    )

    add_heading_num(doc, "6 Выводы")
    add_para(
        doc,
        f"В ходе лабораторной работы построена и выполнена модель простой линейной регрессии "
        f"Salary ~ YearsExperience по методике практикума KNIME (работа № 2). Данные разделены на обучающую "
        f"(20 %, линейная выборка) и тестовую (80 %) части, модель обучена методом наименьших квадратов, "
        f"прогнозы получены для тестовой выборки, качество оценено метриками Numeric Scorer. Достигнуты высокие "
        f"показатели качества (R² = {comma(m['R2'])}; MAPE ≈ {m['MAPE']*100:.1f} %). Цель работы достигнута: "
        f"освоен полный цикл построения, применения и оценки регрессионной модели.",
    )

    add_heading_num(doc, "Список использованных источников")
    for r in [
        "1. Лабораторный практикум KNIME [Электронный ресурс]: методические указания. — Работа № 2: метод линейной регрессии.",
        "2. KNIME AG. KNIME Analytics Platform Documentation [Электронный ресурс]. — URL: https://docs.knime.com (дата обращения: 09.10.2026).",
        "3. Salary Data [Электронный ресурс]. — URL: https://www.kaggle.com/datasets/rkiattisak/salaly-prediction-for-beginer (дата обращения: 09.10.2026).",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        set_run_font(p.add_run(r), size=14)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    build_title_page(doc)
    build_body(doc)

    out = LAB1 / "Отчет_ЛР1_Лаврешин_Линейная_регрессия.docx"
    doc.save(out)
    print("Wrote", out)


if __name__ == "__main__":
    main()
