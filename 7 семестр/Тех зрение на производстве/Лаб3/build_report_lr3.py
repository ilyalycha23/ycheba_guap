# -*- coding: utf-8 -*-
"""ЛР3 Тех зрение — отчёт ГОСТ 7.32 по предобработке изображений (OpenCV)."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

ROOT = Path(Path(r"C:\temp\lab3_tv\lab3_path.txt").read_text(encoding="utf-8").strip())
FIG = ROOT / "figures"
META = json.loads((FIG / "meta.json").read_text(encoding="utf-8"))
FONT = "Times New Roman"


def set_run_font(run, size=14, bold=False, italic=False, name=None):
    name = name or FONT
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_pf(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=1.25, before=0, after=0, line=1.5, keep_with_next=False):
    pf = p.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Cm(first) if first else Cm(0)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if keep_with_next:
        pf.keep_with_next = True


def add_text(doc, text, *, size=14, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             first=1.25, before=0, after=0, keep_with_next=False):
    p = doc.add_paragraph()
    set_pf(p, align=align, first=first, before=before, after=after, keep_with_next=keep_with_next)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic)
    return p


def add_struct(doc, text):
    return add_text(doc, text.upper(), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, after=12, keep_with_next=True)


def add_section(doc, text):
    return add_text(doc, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, before=18, after=12, keep_with_next=True)


def add_caption_fig(doc, text):
    return add_text(doc, text, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, before=6, after=12)


def add_caption_tbl(doc, text):
    return add_text(doc, text, size=12, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, before=10, after=4, keep_with_next=True)


def add_code(doc, text):
    for line in text.split("\n"):
        p = doc.add_paragraph()
        set_pf(p, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, line=1.0)
        set_run_font(p.add_run(line if line else " "), size=10, name="Consolas")


def add_pic(doc, path, width_cm=14.0):
    p = doc.add_paragraph()
    set_pf(p, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, before=8, line=1.0, keep_with_next=True)
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def empty(doc, n=1):
    for _ in range(n):
        p = doc.add_paragraph()
        set_pf(p, first=0, line=1.0)
        set_run_font(p.add_run(""), 14)


def set_borders(table):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def set_borders_nil(table):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def set_cell(cell, text, *, bold=False, size=11):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.0
    pf.first_line_indent = Cm(0)
    set_run_font(p.add_run(str(text)), size=size, bold=bold)


def add_table(doc, headers, rows, widths=None, size=11):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_borders(t)
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, bold=True, size=size)
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            set_cell(t.rows[r].cells[c], val, size=size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def add_toc_line(doc, title, page):
    p = doc.add_paragraph()
    set_pf(p, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, before=2, after=2)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    set_run_font(p.add_run(f"{title}\t{page}"), 14)


def add_biblio(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.left_indent = Cm(0.75)
    pf.first_line_indent = Cm(-0.75)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    set_run_font(p.add_run(text), 14)


def add_page_number(section):
    section.different_first_page_header_footer = True
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_pf(p, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, line=1.0)
    run = p.add_run()
    set_run_font(run, size=14)
    for typ, extra in (("begin", None), ("instr", " PAGE "), ("separate", None), ("end", None)):
        if typ == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = extra
            run._r.append(el)
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), typ)
            run._r.append(el)
    ff = section.first_page_footer
    ff.is_linked_to_previous = False
    if ff.paragraphs:
        ff.paragraphs[0].text = ""


def clear_cell(cell):
    cell.text = ""


def fmt_ms(x):
    return str(x).replace(".", ",")


def build():
    t = META["times_ms"]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    add_page_number(sec)

    # TITLE
    for line in [
        "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ",
        "РОССИЙСКОЙ ФЕДЕРАЦИИ",
        "ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ АВТОНОМНОЕ ОБРАЗОВАТЕЛЬНОЕ",
        "УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ",
        "«САНКТ-ПЕТЕРБУРГСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ",
        "АЭРОКОСМИЧЕСКОГО ПРИБОРОСТРОЕНИЯ»",
    ]:
        add_text(doc, line, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, size=12, after=0)
    add_text(doc, "КАФЕДРА № 42", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, after=16)
    add_text(doc, "ОТЧЕТ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=0)
    add_text(doc, "ЗАЩИЩЕН С ОЦЕНКОЙ", align=WD_ALIGN_PARAGRAPH.CENTER, first=0, after=10)
    add_text(doc, "ПРЕПОДАВАТЕЛЬ", bold=True, first=0, after=4)
    tbl = doc.add_table(rows=2, cols=3)
    set_borders_nil(tbl)
    widths = [Cm(6.5), Cm(5), Cm(5)]
    for ri, row in enumerate([
        ["преподаватель", "", "М.Д. Яушкина"],
        ["должность, уч. степень, звание", "подпись, дата", "инициалы, фамилия"],
    ]):
        for ci, txt in enumerate(row):
            clear_cell(tbl.rows[ri].cells[ci])
            tbl.rows[ri].cells[ci].width = widths[ci]
            p = tbl.rows[ri].cells[ci].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(txt), size=11 if ri == 1 else 14, italic=(ri == 1))
    empty(doc, 2)
    add_text(doc, "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 3", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, after=8)
    add_text(doc, "Предобработка, фильтрация, преобразование", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=0)
    add_text(doc, "изображений и видеокадров на производстве", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, after=8)
    add_text(doc, "по курсу: Системы технического зрения на производстве", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, after=18)
    add_text(doc, "РАБОТУ ВЫПОЛНИЛ", bold=True, first=0, after=4)
    t2 = doc.add_table(rows=2, cols=3)
    set_borders_nil(t2)
    for ri, row in enumerate([
        ["СТУДЕНТ гр. № 4321", "", "И.А. Лаврешин"],
        ["", "подпись, дата", "инициалы, фамилия"],
    ]):
        for ci, txt in enumerate(row):
            clear_cell(t2.rows[ri].cells[ci])
            t2.rows[ri].cells[ci].width = widths[ci]
            p = t2.rows[ri].cells[ci].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(txt), size=11 if ri == 1 else 14, italic=(ri == 1), bold=(ri == 0 and ci == 0))
    empty(doc, 4)
    add_text(doc, "Санкт-Петербург", align=WD_ALIGN_PARAGRAPH.CENTER, first=0)
    add_text(doc, "2026", align=WD_ALIGN_PARAGRAPH.CENTER, first=0)
    doc.add_page_break()

    add_struct(doc, "Содержание")
    for title, page in [
        ("Введение", "3"),
        ("1 Цель работы", "3"),
        ("2 Описание объекта и исходных данных", "3"),
        ("3 Параметры исходного кадра", "4"),
        ("4 Цветовые пространства и гистограмма", "4"),
        ("5 Коррекция яркости и контраста", "5"),
        ("6 Сравнение методов фильтрации", "6"),
        ("7 Изменение размера до 640×640", "7"),
        ("8 Итоговый конвейер и обработка видео", "8"),
        ("9 Примеры исходных и обработанных кадров", "9"),
        ("Заключение", "10"),
        ("Список использованных источников", "10"),
        ("Приложение А Листинг программы", "11"),
    ]:
        add_toc_line(doc, title, page)
    doc.add_page_break()

    add_struct(doc, "Введение")
    add_text(
        doc,
        "В системах технического зрения кадр от камеры является исходным набором данных. "
        "До этапов обнаружения и распознавания изображение необходимо привести к виду, "
        "в котором полезные признаки объекта сохраняются, а влияние шума и неравномерного "
        "освещения уменьшается. В лабораторной работе № 3 выполнен этап предварительной "
        "обработки для выбранного производственного объекта; полученный конвейер будет "
        "использоваться в последующих работах модуля.",
    )

    add_section(doc, "1 Цель работы")
    add_text(
        doc,
        "Цель работы — получить навыки предварительной обработки изображений и видеокадров, "
        "выбора методов фильтрации и преобразования данных для дальнейшего обнаружения, "
        "разметки и распознавания производственных объектов.",
    )

    add_section(doc, "2 Описание объекта и исходных данных")
    add_text(
        doc,
        "В качестве производственного объекта выбрана поверхность горячекатаной стальной "
        "полосы с дефектами. Использован датасет NEU-DET (Northeastern University / Kaggle), "
        "который указан в методических указаниях как пример исходных данных. Классы дефектов: "
        "scratches, patches, inclusion, crazing, pitted_surface, rolled-in_scale. "
        "Объект подходит для последующих работ по обнаружению и классификации дефектов.",
    )
    add_text(
        doc,
        f"Отобрано {META['n_images']} исходных кадров (по 5 из каждого класса) и собран "
        f"видеоролик {META['video']} длительностью {META['video_duration_s']} с при частоте "
        f"{META['fps']} кадр/с. Рабочий кадр для анализа приведён к разрешению "
        f"{META['resolution']} (типичный видеопоток стенда). Среда: Python, OpenCV.",
    )
    add_pic(doc, FIG / "00_classes.jpg", 14.0)
    add_caption_fig(doc, "Рисунок 1 − Примеры классов дефектов NEU-DET")
    add_caption_tbl(doc, "Таблица 1 − Структура каталогов проекта")
    add_table(
        doc,
        ["Каталог", "Назначение"],
        [
            ["data/raw/", "исходные изображения и видео"],
            ["data/frames/", "кадры, извлечённые из видео"],
            ["data/processed/", "итоговые обработанные изображения"],
            ["results/", "иллюстрации и графики для отчёта"],
        ],
        widths=[5, 11.5],
    )

    add_section(doc, "3 Параметры исходного кадра")
    add_text(
        doc,
        f"Рабочий кадр имеет разрешение {META['resolution']} пикселей, число каналов — "
        f"{META['channels']} (BGR, uint8). Для видеоролика частота кадров составляет "
        f"{META['fps']} кадр/с, что задаёт бюджет времени на обработку одного кадра "
        "примерно 40 мс при работе без пропуска кадров.",
    )
    add_pic(doc, FIG / "01_original.jpg", 13.5)
    add_caption_fig(doc, "Рисунок 2 − Рабочий исходный кадр (дефект типа scratch)")

    add_section(doc, "4 Цветовые пространства и гистограмма")
    add_text(
        doc,
        "Исходный кадр преобразован в оттенки серого и в пространство HSV. Для выбранного "
        "объекта основными признаками являются локальный контраст дефекта относительно "
        "фоновой текстуры металла и геометрия дефекта. Цвет поверхности не является "
        "ключевым диагностическим признаком, поэтому основные операции коррекции выполняются "
        "по каналу яркости; цветное изображение сохраняется для последующих работ.",
    )
    add_pic(doc, FIG / "06_colorspaces.png", 15.5)
    add_caption_fig(doc, "Рисунок 3 − Преобразование цветового пространства")
    add_text(
        doc,
        f"По гистограмме яркости среднее значение составляет {fmt_ms(META['gray_mean'])}, "
        f"среднеквадратическое отклонение — {fmt_ms(META['gray_std'])}. Распределение "
        "отражает текстуру металла и область дефекта; контраст дефекта относительно фона "
        "достаточен для дальнейшей обработки.",
    )
    add_pic(doc, FIG / "07_hist_original.png", 12.5)
    add_caption_fig(doc, "Рисунок 4 − Гистограмма яркости исходного кадра")

    add_section(doc, "5 Коррекция яркости и контраста")
    add_text(
        doc,
        "Сравнены три способа коррекции: линейное изменение яркости и контраста "
        "(α = 1,15; β = 10), гамма-коррекция с γ = 0,8 и локальное выравнивание "
        "гистограммы CLAHE (clipLimit = 2,0; сетка 8×8) в пространстве LAB по каналу L.",
    )
    add_pic(doc, FIG / "11_corrections.png", 14.5)
    add_caption_fig(doc, "Рисунок 5 − Сравнение способов коррекции яркости и контраста")
    add_text(
        doc,
        f"После CLAHE среднее значение яркости увеличилось до {fmt_ms(META['clahe_mean'])}, "
        f"СКО — до {fmt_ms(META['clahe_std'])}. Локальный контраст дефектов относительно "
        "фона улучшился без глобального «выбеливания» всего кадра. Для итогового конвейера "
        "выбран CLAHE.",
    )

    add_section(doc, "6 Сравнение методов фильтрации")
    add_text(
        doc,
        "Для сравнения гауссова, медианного и билатерального фильтров на кадр после CLAHE "
        "добавлен контрольный импульсный шум (только для эксперимента; в итоговый набор "
        "данных искусственный шум не входит). Измерение времени — не менее 30 прогонов "
        f"на кадре {META['resolution']}.",
    )
    add_pic(doc, FIG / "16_filters.png", 14.5)
    add_caption_fig(doc, "Рисунок 6 − Сравнение результатов фильтрации")
    add_caption_tbl(doc, "Таблица 2 − Среднее время фильтрации")
    add_table(
        doc,
        ["Метод", "Параметры", "Среднее время, мс"],
        [
            ["Гауссов фильтр", "ядро 5×5", fmt_ms(t["gaussian_5x5"])],
            ["Медианный фильтр", "ядро 5×5", fmt_ms(t["median_5x5"])],
            ["Билатеральный фильтр", "d = 9; σ = 75", fmt_ms(t["bilateral_9_75"])],
        ],
        widths=[5.5, 5.5, 5.5],
    )
    add_text(
        doc,
        "Гауссов фильтр быстрее всех, но сильнее размывает границы дефектов. "
        "Билатеральный лучше сохраняет контуры, однако занимает существенную долю "
        "бюджета 40 мс на кадр. Медианный фильтр эффективно подавляет точечный шум "
        "и сохраняет границы при приемлемом времени. Для итогового конвейера выбран "
        "медианный фильтр 5×5.",
    )

    add_section(doc, "7 Изменение размера до 640×640")
    add_text(
        doc,
        "Кадр приведён к размеру 640×640 пикселей с сохранением соотношения сторон: "
        "1280×720 масштабируется до 640×360, свободные области сверху и снизу заполняются "
        "значением 114. Геометрия дефектов не растягивается по одной из осей.",
    )
    add_pic(doc, FIG / "18_resize.png", 14.5)
    add_caption_fig(doc, "Рисунок 7 − Результат масштабирования с дополнением до 640×640")

    add_section(doc, "8 Итоговый конвейер и обработка видео")
    add_text(
        doc,
        "Из видеоролика извлечено не менее 50 кадров через равные интервалы "
        f"(сохранено {META['frames_saved']} кадров). К каждому кадру применён единый конвейер.",
    )
    add_caption_tbl(doc, "Таблица 3 − Параметры итогового конвейера предобработки")
    add_table(
        doc,
        ["Этап", "Параметры", "Назначение"],
        [
            ["1. Исходный кадр", f"{META['resolution']}, BGR", "Получение данных"],
            ["2. Коррекция контраста", "CLAHE: 2,0; 8×8", "Неравномерное освещение"],
            ["3. Фильтрация", "медианный 5×5", "Точечный шум"],
            ["4. Масштабирование", "→ 640×360", "Снижение нагрузки"],
            ["5. Дополнение", "до 640×640, value=114", "Единый размер без искажения"],
        ],
        widths=[4.5, 5.0, 7.0],
        size=10,
    )
    add_text(
        doc,
        f"Среднее время полного конвейера на одном кадре {META['resolution']} составляет "
        f"{fmt_ms(META['pipeline_ms'])} мс, что укладывается в бюджет ~40 мс при 25 кадр/с "
        "и оставляет запас на последующие этапы обнаружения.",
    )

    add_section(doc, "9 Примеры исходных и обработанных кадров")
    add_text(
        doc,
        "Для проверки устойчивости конвейера обработаны кадры с различным положением "
        "дефектов и уровнем освещения. Параметры CLAHE, медианного фильтра и масштабирования "
        "одинаковы для всех кадров.",
    )
    for i in range(1, 4):
        add_pic(doc, FIG / f"19_pair_{i}.jpg", 15.0)
        add_caption_fig(doc, f"Рисунок {7 + i} − Пара «исходный кадр – обработанный кадр» № {i}")

    add_struct(doc, "Заключение")
    add_text(
        doc,
        "Сформирован минимальный конвейер предобработки кадров NEU-DET (дефекты поверхности "
        "стальной полосы): CLAHE → медианная фильтрация 5×5 → масштабирование с сохранением "
        "соотношения сторон → дополнение до 640×640. Цветовая информация для данного объекта "
        "не является основным признаком; цветные кадры сохранены для следующих лабораторных "
        f"работ. Каталог data/processed содержит {META['frames_saved']} обработанных кадров.",
    )
    add_text(
        doc,
        "В последующих работах модуля будут использоваться полученные обработанные изображения "
        "для обнаружения признаков, разметки и обучения модели без повторной настройки "
        "параметров предобработки.",
    )

    add_struct(doc, "Список использованных источников")
    add_biblio(
        doc,
        "1. Предобработка, фильтрация, преобразование изображений и видеокадров на производстве : "
        "метод. указания к лабораторной работе № 3. – Санкт-Петербург : ГУАП, 2026.",
    )
    add_biblio(
        doc,
        "2. OpenCV Documentation. Image Processing [Электронный ресурс]. – Режим доступа: "
        "https://docs.opencv.org/ (дата обращения: 06.10.2026).",
    )
    add_biblio(
        doc,
        "3. Song K., Yan Y. A noise robust method based on completed local binary patterns for "
        "hot-rolled steel strip surface defects // Applied Surface Science. – 2013. – Vol. 285. – P. 858–864.",
    )

    doc.add_page_break()
    add_struct(doc, "Приложение А")
    add_text(doc, "Листинг программы", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, after=8)
    add_text(
        doc,
        "Фрагмент файла project/run_lab3.py: коррекция CLAHE, сравнение фильтров, "
        "масштабирование с padding и применение конвейера к кадрам видео.",
    )
    add_caption_tbl(doc, "Листинг А.1 − Ключевые операции предобработки")
    add_code(
        doc,
        "lab = cv2.cvtColor(img, cv2.COLOR_BGR2LAB)\nl, a, b = cv2.split(lab)\n"
        "clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))\n"
        "l2 = clahe.apply(l)\nclahe_img = cv2.cvtColor(cv2.merge([l2, a, b]), cv2.COLOR_LAB2BGR)\n\n"
        "gaussian = cv2.GaussianBlur(frame, (5, 5), 0)\n"
        "median = cv2.medianBlur(frame, 5)\n"
        "bilateral = cv2.bilateralFilter(frame, 9, 75, 75)\n\n"
        "def resize_with_padding(image, size=640, value=114):\n"
        "    h, w = image.shape[:2]\n"
        "    scale = min(size / w, size / h)\n"
        "    nw, nh = int(w * scale), int(h * scale)\n"
        "    resized = cv2.resize(image, (nw, nh), interpolation=cv2.INTER_AREA)\n"
        "    top = (size - nh) // 2; bottom = size - nh - top\n"
        "    left = (size - nw) // 2; right = size - nw - left\n"
        "    return cv2.copyMakeBorder(resized, top, bottom, left, right,\n"
        "        cv2.BORDER_CONSTANT, value=(value, value, value))",
    )

    out = ROOT / "ЛР3_ТЕХ_Лаврешин.docx"
    try:
        doc.save(out)
    except PermissionError:
        out = ROOT / "ЛР3_ТЕХ_Лаврешин_новый.docx"
        doc.save(out)
    print("WROTE", out)
    return out


if __name__ == "__main__":
    build()
